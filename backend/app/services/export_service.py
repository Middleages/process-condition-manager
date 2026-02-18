import io
import logging
import re
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models import (
    ExportSystem, ExportColumnMapping, Project, ProjectLayer,
    EquipmentAssignment,
)

logger = logging.getLogger(__name__)


def sanitize_filename(name: str) -> str:
    """Replace special characters in filenames with underscores."""
    return re.sub(r'[^\w\-.]', '_', name)


class ExportService:
    """Generates Excel exports for approved project condition tables."""

    # --- Public API ---

    async def get_systems(self, db: AsyncSession) -> list[dict]:
        """Return all active export systems with column mapping counts."""
        query = (
            select(ExportSystem)
            .where(ExportSystem.is_active.is_(True))
            .options(selectinload(ExportSystem.column_mappings))
            .order_by(ExportSystem.system_name)
        )
        result = await db.execute(query)
        systems = result.scalars().all()
        return [
            {
                "id": s.id,
                "system_name": s.system_name,
                "format_type": s.format_type,
                "description": s.description,
                "column_count": len(s.column_mappings),
                "is_active": s.is_active,
            }
            for s in systems
        ]

    async def generate(
        self, db: AsyncSession, project_id: int, system_id: int
    ) -> tuple[bytes, str]:
        """Generate Excel bytes for a single system export.
        Returns (excel_bytes, filename).
        """
        project = await self._get_project(db, project_id)
        system = await self._get_export_system(db, system_id)
        mappings = await self._get_column_mappings(db, system_id)
        layers_data = await self._get_project_layers(db, project_id)
        product_name = project.product.product_name

        if system.format_type == "TYPE_A":
            excel_bytes = self._generate_type_a(product_name, layers_data, mappings)
        elif system.format_type == "TYPE_B":
            equipment = await self._get_equipment_assignments(db, [pl.id for pl in layers_data])
            excel_bytes = self._generate_type_b(product_name, layers_data, mappings, equipment)
        elif system.format_type == "TYPE_C":
            unit_mappings = (system.additional_config or {}).get("unit_mappings", {})
            excel_bytes = self._generate_type_c(product_name, layers_data, mappings, unit_mappings)
        else:
            raise ValueError(f"Unknown format type: {system.format_type}")

        filename = f"{sanitize_filename(product_name)}_{sanitize_filename(system.system_name)}.xlsx"
        return excel_bytes, filename

    async def generate_preview(
        self, db: AsyncSession, project_id: int, system_id: int, limit: int = 5
    ) -> dict:
        """Generate preview data (first N rows) as JSON."""
        project = await self._get_project(db, project_id)
        system = await self._get_export_system(db, system_id)
        mappings = await self._get_column_mappings(db, system_id)
        layers_data = await self._get_project_layers(db, project_id)
        product_name = project.product.product_name

        if system.format_type == "TYPE_A":
            headers, all_rows = self._build_type_a_data(product_name, layers_data, mappings)
        elif system.format_type == "TYPE_B":
            equipment = await self._get_equipment_assignments(db, [pl.id for pl in layers_data])
            headers, all_rows = self._build_type_b_data(product_name, layers_data, mappings, equipment)
        elif system.format_type == "TYPE_C":
            unit_mappings = (system.additional_config or {}).get("unit_mappings", {})
            headers, all_rows = self._build_type_c_data(product_name, layers_data, mappings, unit_mappings)
        else:
            raise ValueError(f"Unknown format type: {system.format_type}")

        return {
            "system_name": system.system_name,
            "format_type": system.format_type,
            "headers": headers,
            "rows": all_rows[:limit],
            "total_rows": len(all_rows),
        }

    async def generate_bulk(
        self, db: AsyncSession, project_id: int, system_ids: list[int]
    ) -> tuple[bytes, str]:
        """Generate ZIP file containing Excel files for multiple systems.
        Returns (zip_bytes, filename).
        """
        project = await self._get_project(db, project_id)
        product_name = project.product.product_name

        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
            for sid in system_ids:
                excel_bytes, fname = await self.generate(db, project_id, sid)
                zf.writestr(fname, excel_bytes)

        from datetime import datetime
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        zip_filename = f"{sanitize_filename(product_name)}_export_{timestamp}.zip"
        return zip_buffer.getvalue(), zip_filename

    # --- Data Retrieval Helpers ---

    async def _get_project(self, db: AsyncSession, project_id: int) -> Project:
        """Fetch project with product info."""
        query = (
            select(Project)
            .where(Project.id == project_id)
            .options(selectinload(Project.product))
        )
        result = await db.execute(query)
        project = result.scalar_one_or_none()
        if not project:
            raise ValueError(f"Project {project_id} not found")
        return project

    async def _get_export_system(self, db: AsyncSession, system_id: int) -> ExportSystem:
        """Fetch export system by ID."""
        query = select(ExportSystem).where(ExportSystem.id == system_id)
        result = await db.execute(query)
        system = result.scalar_one_or_none()
        if not system:
            raise ValueError(f"Export system {system_id} not found")
        return system

    async def _get_column_mappings(
        self, db: AsyncSession, system_id: int
    ) -> list[ExportColumnMapping]:
        """Fetch column mappings for a system, eagerly loading column_definition."""
        query = (
            select(ExportColumnMapping)
            .where(ExportColumnMapping.export_system_id == system_id)
            .options(selectinload(ExportColumnMapping.column_definition))
            .order_by(ExportColumnMapping.sort_order)
        )
        result = await db.execute(query)
        return list(result.scalars().all())

    async def _get_project_layers(
        self, db: AsyncSession, project_id: int
    ) -> list[ProjectLayer]:
        """Fetch project layers with layer info, ordered by sort_order."""
        query = (
            select(ProjectLayer)
            .where(ProjectLayer.project_id == project_id)
            .options(selectinload(ProjectLayer.layer))
            .order_by(ProjectLayer.sort_order)
        )
        result = await db.execute(query)
        return list(result.scalars().all())

    async def _get_equipment_assignments(
        self, db: AsyncSession, project_layer_ids: list[int]
    ) -> dict[int, list[EquipmentAssignment]]:
        """Fetch equipment assignments grouped by project_layer_id."""
        if not project_layer_ids:
            return {}
        query = (
            select(EquipmentAssignment)
            .where(EquipmentAssignment.project_layer_id.in_(project_layer_ids))
            .order_by(EquipmentAssignment.project_layer_id, EquipmentAssignment.sort_order)
        )
        result = await db.execute(query)
        assignments = result.scalars().all()

        grouped: dict[int, list[EquipmentAssignment]] = {}
        for a in assignments:
            grouped.setdefault(a.project_layer_id, []).append(a)
        return grouped

    # --- Type A: Horizontal (1 row per layer) ---

    def _build_type_a_data(
        self, product_name: str, layers: list[ProjectLayer],
        mappings: list[ExportColumnMapping]
    ) -> tuple[list[str], list[dict]]:
        """Build Type A tabular data."""
        headers = ["LAYER_ID", "PRODUCT_ID"]
        for m in mappings:
            headers.append(m.target_column_name)

        rows = []
        for pl in layers:
            row = {
                "LAYER_ID": pl.layer.layer_name,
                "PRODUCT_ID": product_name,
            }
            conditions = pl.conditions or {}
            for m in mappings:
                col_name = m.column_definition.column_name
                value = conditions.get(col_name)
                if value is None and m.is_required:
                    logger.warning(
                        "Required column %s is null for layer %s",
                        col_name, pl.layer.layer_name
                    )
                row[m.target_column_name] = value if value is not None else ""
            rows.append(row)
        return headers, rows

    def _generate_type_a(
        self, product_name: str, layers: list[ProjectLayer],
        mappings: list[ExportColumnMapping]
    ) -> bytes:
        """Type A: 1 row per layer, horizontal format."""
        headers, rows = self._build_type_a_data(product_name, layers, mappings)
        return self._write_excel(headers, rows)

    # --- Type B: Equipment-Split Multi-Row ---

    def _build_type_b_data(
        self, product_name: str, layers: list[ProjectLayer],
        mappings: list[ExportColumnMapping],
        equipment: dict[int, list[EquipmentAssignment]]
    ) -> tuple[list[str], list[dict]]:
        """Build Type B tabular data."""
        headers = ["LAYER_ID", "PRODUCT_ID", "EQUIP_ID"]
        for m in mappings:
            headers.append(m.target_column_name)

        rows = []
        for pl in layers:
            equip_list = equipment.get(pl.id, [])
            conditions = pl.conditions or {}

            if not equip_list:
                # No equipment -> single row with empty EQUIP_ID
                row = {
                    "LAYER_ID": pl.layer.layer_name,
                    "PRODUCT_ID": product_name,
                    "EQUIP_ID": "",
                }
                for m in mappings:
                    col_name = m.column_definition.column_name
                    row[m.target_column_name] = conditions.get(col_name, "")
                rows.append(row)
            else:
                for ea in equip_list:
                    row = {
                        "LAYER_ID": pl.layer.layer_name,
                        "PRODUCT_ID": product_name,
                        "EQUIP_ID": ea.equipment_id,
                    }
                    overrides = ea.equipment_params or {}
                    for m in mappings:
                        col_name = m.column_definition.column_name
                        # Equipment override takes precedence
                        if col_name in overrides:
                            row[m.target_column_name] = overrides[col_name]
                        else:
                            row[m.target_column_name] = conditions.get(col_name, "")
                    rows.append(row)
        return headers, rows

    def _generate_type_b(
        self, product_name: str, layers: list[ProjectLayer],
        mappings: list[ExportColumnMapping],
        equipment: dict[int, list[EquipmentAssignment]]
    ) -> bytes:
        """Type B: Equipment-split multi-row format."""
        headers, rows = self._build_type_b_data(product_name, layers, mappings, equipment)
        return self._write_excel(headers, rows)

    # --- Type C: Key-Value Vertical Transpose ---

    def _build_type_c_data(
        self, product_name: str, layers: list[ProjectLayer],
        mappings: list[ExportColumnMapping],
        unit_mappings: dict[str, str]
    ) -> tuple[list[str], list[dict]]:
        """Build Type C tabular data."""
        headers = ["LAYER_ID", "PRODUCT_ID", "PARAM_KEY", "PARAM_VALUE", "UNIT"]

        rows = []
        for pl in layers:
            conditions = pl.conditions or {}
            for m in mappings:
                col_name = m.column_definition.column_name
                value = conditions.get(col_name, "")
                unit = unit_mappings.get(col_name, "")
                rows.append({
                    "LAYER_ID": pl.layer.layer_name,
                    "PRODUCT_ID": product_name,
                    "PARAM_KEY": m.target_column_name,
                    "PARAM_VALUE": value,
                    "UNIT": unit,
                })
        return headers, rows

    def _generate_type_c(
        self, product_name: str, layers: list[ProjectLayer],
        mappings: list[ExportColumnMapping],
        unit_mappings: dict[str, str]
    ) -> bytes:
        """Type C: Key-value vertical transpose format."""
        headers, rows = self._build_type_c_data(product_name, layers, mappings, unit_mappings)
        return self._write_excel(headers, rows)

    # --- Excel Writing ---

    def _write_excel(self, headers: list[str], rows: list[dict]) -> bytes:
        """Write data to Excel and return bytes."""
        wb = Workbook()
        ws = wb.active

        # Write headers with bold font
        bold_font = Font(bold=True)
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = bold_font

        # Write data rows
        for row_idx, row_data in enumerate(rows, 2):
            for col_idx, header in enumerate(headers, 1):
                value = row_data.get(header, "")
                ws.cell(row=row_idx, column=col_idx, value=value)

        # Auto-adjust column widths
        for col_idx, header in enumerate(headers, 1):
            max_len = len(str(header))
            for row_idx in range(2, len(rows) + 2):
                cell_value = ws.cell(row=row_idx, column=col_idx).value
                if cell_value is not None:
                    max_len = max(max_len, len(str(cell_value)))
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 2, 50)

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
