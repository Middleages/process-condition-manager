import io
import logging
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from app.models import ProjectLayer, ExportColumnMapping

logger = logging.getLogger(__name__)


def sanitize_filename(name: str) -> str:
    import re
    return re.sub(r'[^\w\-.]', '_', name)


def _get_mapping_value(
    mapping: ExportColumnMapping,
    conditions: dict,
    external_data: dict[int, dict[str, dict[str, Any]]],
    layer_id: int,
) -> Any:
    """Get the value for a mapping from either conditions or external data.

    For source_type='external': looks up external_data[data_source_id][str(layer_id)][source_column_name].
    For source_type='condition': looks up conditions[column_definition.column_name].
    Returns empty string when no value is found.
    """
    if mapping.source_type == "external":
        ds_id = mapping.data_source_id
        if ds_id is not None and ds_id in external_data:
            layer_data = external_data[ds_id].get(str(layer_id), {})
            return layer_data.get(mapping.source_column_name, "")
        return ""
    else:
        # condition type - existing logic
        if mapping.column_definition is None:
            return ""
        return conditions.get(mapping.column_definition.column_name, "")


def build_type_a_data(
    product_name: str, layers: list[ProjectLayer],
    mappings: list[ExportColumnMapping],
    external_data: dict[int, dict[str, dict[str, Any]]] | None = None,
) -> tuple[list[str], list[dict]]:
    ext_data = external_data or {}
    headers = ["LAYER_ID", "PRODUCT_ID"]
    for m in mappings:
        headers.append(m.target_column_name)

    rows = []
    for pl in layers:
        row = {
            "LAYER_ID": pl.layer_name,
            "PRODUCT_ID": product_name,
        }
        conditions = pl.conditions or {}
        for m in mappings:
            value = _get_mapping_value(m, conditions, ext_data, pl.id)
            if value in (None, "") and m.is_required:
                if m.source_type == "condition" and m.column_definition:
                    logger.warning(
                        "Required column %s is null for layer %s",
                        m.column_definition.column_name, pl.layer_name
                    )
                elif m.source_type == "external":
                    logger.warning(
                        "Required external column %s is null for layer %s",
                        m.source_column_name, pl.layer_name
                    )
            row[m.target_column_name] = value if value is not None else ""
        rows.append(row)
    return headers, rows


def generate_type_a(
    product_name: str, layers: list[ProjectLayer],
    mappings: list[ExportColumnMapping],
    external_data: dict[int, dict[str, dict[str, Any]]] | None = None,
) -> bytes:
    headers, rows = build_type_a_data(product_name, layers, mappings, external_data)
    return write_excel(headers, rows)


def build_type_b_data(
    product_name: str, layers: list[ProjectLayer],
    mappings: list[ExportColumnMapping],
    system_config: dict | None = None,
    external_data: dict[int, dict[str, dict[str, Any]]] | None = None,
) -> tuple[list[str], list[dict]]:
    """conditions JSONB의 EQP 컬럼(EQP_01..EQP_20)을 읽어 레이어당 설비별 행을 생성한다.

    system_config.equip_vary_mapping: 설비별로 값이 다른 컬럼과 EQP 파라미터 접미사 매핑.
    예) {"SC_EXPOSE_ENERGY_mJ": "_ET", "SC_EXPOSE_FOCUS_um": "_FOCUS"}
    """
    ext_data = external_data or {}
    # equip_vary_mapping: {조건 컬럼명 -> EQP 파라미터 접미사} 형식
    equip_vary_mapping: dict[str, str] = (system_config or {}).get("equip_vary_mapping", {})

    headers = ["LAYER_ID", "PRODUCT_ID", "EQUIP_ID"]
    for m in mappings:
        headers.append(m.target_column_name)

    rows = []
    for pl in layers:
        conditions = pl.conditions or {}
        found_equipment = False

        # EQP_01 부터 EQP_20까지 슬롯을 순회하여 설비별 행 생성
        for slot in range(1, 21):
            nn = f"{slot:02d}"
            eqp_name = conditions.get(f"EQP_{nn}", "")
            if not eqp_name:
                # 빈 슬롯은 건너뜀 (R7.4: 이름 없이 파라미터만 있으면 무시)
                continue

            found_equipment = True
            row = {
                "LAYER_ID": pl.layer_name,
                "PRODUCT_ID": product_name,
                "EQUIP_ID": eqp_name,
            }

            for m in mappings:
                col_name = m.column_definition.column_name if m.column_definition else None
                if col_name and col_name in equip_vary_mapping:
                    # 설비별로 값이 다른 컬럼: EQP_{NN}{접미사} 컬럼에서 읽음
                    suffix = equip_vary_mapping[col_name]
                    row[m.target_column_name] = conditions.get(f"EQP_{nn}{suffix}", "")
                else:
                    # 공통 컬럼: 기본 conditions에서 읽음
                    row[m.target_column_name] = _get_mapping_value(
                        m, conditions, ext_data, pl.id
                    )
            rows.append(row)

        if not found_equipment:
            # 설비가 없는 레이어는 빈 EQUIP_ID로 단일 행 생성 (기존 fallback 동작 유지)
            row = {
                "LAYER_ID": pl.layer_name,
                "PRODUCT_ID": product_name,
                "EQUIP_ID": "",
            }
            for m in mappings:
                row[m.target_column_name] = _get_mapping_value(m, conditions, ext_data, pl.id)
            rows.append(row)

    return headers, rows


def generate_type_b(
    product_name: str, layers: list[ProjectLayer],
    mappings: list[ExportColumnMapping],
    system_config: dict | None = None,
    external_data: dict[int, dict[str, dict[str, Any]]] | None = None,
) -> bytes:
    """Type B Excel 파일을 생성한다. EQP 컬럼 기반 설비 분할 방식 사용."""
    headers, rows = build_type_b_data(product_name, layers, mappings, system_config, external_data)
    return write_excel(headers, rows)


def build_type_c_data(
    product_name: str, layers: list[ProjectLayer],
    mappings: list[ExportColumnMapping],
    unit_mappings: dict[str, str],
    external_data: dict[int, dict[str, dict[str, Any]]] | None = None,
) -> tuple[list[str], list[dict]]:
    ext_data = external_data or {}
    headers = ["LAYER_ID", "PRODUCT_ID", "PARAM_KEY", "PARAM_VALUE", "UNIT"]

    rows = []
    for pl in layers:
        conditions = pl.conditions or {}
        for m in mappings:
            value = _get_mapping_value(m, conditions, ext_data, pl.id)
            # unit_mappings keyed by condition column name (only meaningful for condition type)
            if m.source_type == "condition" and m.column_definition:
                unit = unit_mappings.get(m.column_definition.column_name, "")
            else:
                unit = ""
            rows.append({
                "LAYER_ID": pl.layer_name,
                "PRODUCT_ID": product_name,
                "PARAM_KEY": m.target_column_name,
                "PARAM_VALUE": value if value is not None else "",
                "UNIT": unit,
            })
    return headers, rows


def generate_type_c(
    product_name: str, layers: list[ProjectLayer],
    mappings: list[ExportColumnMapping],
    unit_mappings: dict[str, str],
    external_data: dict[int, dict[str, dict[str, Any]]] | None = None,
) -> bytes:
    headers, rows = build_type_c_data(product_name, layers, mappings, unit_mappings, external_data)
    return write_excel(headers, rows)


def write_excel(headers: list[str], rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active

    # Write headers with bold font
    bold_font = Font(bold=True)
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = bold_font

    # Write data rows
    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, header in enumerate(headers, 1):
            value = row_data.get(header, "")
            ws.cell(row=row_idx, column=col_idx, value=value)

    # Auto-adjust column widths
    for col_idx, header in enumerate(headers, 1):
        max_len = len(str(header))
        for row_idx in range(2, len(rows) + 2):
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            if cell_value is not None:
                max_len = max(max_len, len(str(cell_value)))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 2, 50)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
