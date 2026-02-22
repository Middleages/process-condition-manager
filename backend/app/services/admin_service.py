"""Admin service for XML mapping and validation rule management."""

import json
from datetime import datetime
from typing import BinaryIO
from sqlalchemy import select, delete
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload
from fastapi import HTTPException
from openpyxl import load_workbook

from app.models import (
    RecipeXmlMapping, ColumnDefinition, ColumnCategory, ColumnValidation,
    ChangeLog, ProjectLayer, Project, Product, Layer, User,
)
from app.schemas.admin import (
    RecipeMappingResponse, RecipeMappingCreate, RecipeMappingUpdate,
    ValidationRuleCreate, BulkUploadResponse,
    ColumnSelectOptionsResponse, SelectOptionsUpdate,
    AuditLogEntry, AuditLogListResponse,
)

from app.constants import ALLOWED_VALUE_TRANSFORMS, ALLOWED_RULE_TYPES


# ---------------------------------------------------------------------------
# Recipe XML Mapping Functions
# ---------------------------------------------------------------------------

async def _build_mapping_response(
    db: AsyncSession,
    mapping_id: int,
) -> RecipeMappingResponse:
    """Build a RecipeMappingResponse by loading the mapping with related data."""
    result = await db.execute(
        select(RecipeXmlMapping, ColumnDefinition, ColumnCategory)
        .join(ColumnDefinition, RecipeXmlMapping.column_id == ColumnDefinition.id)
        .join(ColumnCategory, ColumnDefinition.category_id == ColumnCategory.id)
        .where(RecipeXmlMapping.id == mapping_id)
    )
    mapping, col_def, category = result.one()

    return RecipeMappingResponse(
        id=mapping.id,
        xpath=mapping.xpath,
        column_id=mapping.column_id,
        column_name=col_def.column_name,
        display_name=col_def.display_name,
        category_code=category.category_code,
        data_type=col_def.data_type,
        value_transform=mapping.value_transform,
        is_active=mapping.is_active,
        created_at=mapping.created_at,
    )


async def list_mappings(
    db: AsyncSession,
    is_active: bool | None = None,
    search: str | None = None,
) -> list[RecipeMappingResponse]:
    """List recipe XML mappings with optional filters."""
    query = (
        select(RecipeXmlMapping, ColumnDefinition, ColumnCategory)
        .join(ColumnDefinition, RecipeXmlMapping.column_id == ColumnDefinition.id)
        .join(ColumnCategory, ColumnDefinition.category_id == ColumnCategory.id)
    )

    # Apply filters
    if is_active is not None:
        query = query.where(RecipeXmlMapping.is_active == is_active)

    if search:
        # Search in xpath, column_name, or display_name
        search_pattern = f"%{search}%"
        query = query.where(
            (RecipeXmlMapping.xpath.ilike(search_pattern))
            | (ColumnDefinition.column_name.ilike(search_pattern))
            | (ColumnDefinition.display_name.ilike(search_pattern))
        )

    result = await db.execute(query)
    rows = result.all()

    mappings = []
    for mapping, col_def, category in rows:
        mappings.append(
            RecipeMappingResponse(
                id=mapping.id,
                xpath=mapping.xpath,
                column_id=mapping.column_id,
                column_name=col_def.column_name,
                display_name=col_def.display_name,
                category_code=category.category_code,
                data_type=col_def.data_type,
                value_transform=mapping.value_transform,
                is_active=mapping.is_active,
                created_at=mapping.created_at,
            )
        )

    return mappings


async def create_mapping(
    db: AsyncSession,
    data: RecipeMappingCreate,
) -> RecipeMappingResponse:
    """Create a new recipe XML mapping."""
    # Validate column exists
    col_def = await db.get(ColumnDefinition, data.column_id)
    if not col_def:
        raise HTTPException(status_code=404, detail="Column not found")

    # Validate value_transform
    if data.value_transform is not None and data.value_transform not in ALLOWED_VALUE_TRANSFORMS:
        raise HTTPException(status_code=422, detail=f"Invalid value_transform: {data.value_transform}")

    # Create mapping
    mapping = RecipeXmlMapping(
        xpath=data.xpath,
        column_id=data.column_id,
        value_transform=data.value_transform,
        is_active=True,
    )
    db.add(mapping)
    await db.commit()
    await db.refresh(mapping)

    # Build and return response
    return await _build_mapping_response(db, mapping.id)


async def update_mapping(
    db: AsyncSession,
    mapping_id: int,
    data: RecipeMappingUpdate,
) -> RecipeMappingResponse:
    """Update an existing recipe XML mapping."""
    # Find mapping
    mapping = await db.get(RecipeXmlMapping, mapping_id)
    if not mapping:
        raise HTTPException(status_code=404, detail="Mapping not found")

    # Validate value_transform if provided
    if data.value_transform is not None and data.value_transform not in ALLOWED_VALUE_TRANSFORMS:
        raise HTTPException(status_code=422, detail=f"Invalid value_transform: {data.value_transform}")

    # Update fields
    if data.xpath is not None:
        mapping.xpath = data.xpath
    if data.column_id is not None:
        # Validate new column exists
        col_def = await db.get(ColumnDefinition, data.column_id)
        if not col_def:
            raise HTTPException(status_code=404, detail="Column not found")
        mapping.column_id = data.column_id
    if data.value_transform is not None:
        mapping.value_transform = data.value_transform
    if data.is_active is not None:
        mapping.is_active = data.is_active

    await db.commit()
    await db.refresh(mapping)

    # Build and return response
    return await _build_mapping_response(db, mapping.id)


async def delete_mapping(
    db: AsyncSession,
    mapping_id: int,
) -> None:
    """Delete a recipe XML mapping."""
    mapping = await db.get(RecipeXmlMapping, mapping_id)
    if not mapping:
        raise HTTPException(status_code=404, detail="Mapping not found")

    await db.delete(mapping)
    await db.commit()


# ---------------------------------------------------------------------------
# Validation Rule Functions
# ---------------------------------------------------------------------------

async def list_columns_with_validations(
    db: AsyncSession,
    category_code: str | None = None,
) -> list:
    """List columns grouped by category with their validation rules.

    Returns the same ColumnCategory structure as GET /api/columns
    to maintain frontend compatibility.
    """
    query = (
        select(ColumnCategory)
        .options(
            selectinload(ColumnCategory.columns)
            .selectinload(ColumnDefinition.validations)
        )
        .order_by(ColumnCategory.sort_order)
    )

    if category_code:
        query = query.where(ColumnCategory.category_code == category_code)

    result = await db.execute(query)
    categories = result.scalars().unique().all()

    # Sort columns within each category
    for cat in categories:
        cat.columns.sort(key=lambda c: c.sort_order)

    return categories


def _validate_cross_layer_rule_config(rule_config: dict) -> None:
    """cross_layer 규칙의 rule_config 유효성을 검증한다.

    check_type별 필수 필드와 허용 값을 확인하고,
    유효하지 않으면 HTTPException(422)을 발생시킨다.
    """
    from app.services.cross_layer_validation_service import (
        ALLOWED_CHECK_TYPES,
        ALLOWED_OPERATORS,
        ALLOWED_COMPATIBILITY_TYPES,
    )

    check_type = rule_config.get("check_type")
    if not check_type or check_type not in ALLOWED_CHECK_TYPES:
        raise HTTPException(
            status_code=422,
            detail=(
                f"cross_layer 규칙의 check_type은 다음 중 하나여야 합니다: "
                f"{', '.join(ALLOWED_CHECK_TYPES)}. 현재 값: '{check_type}'"
            ),
        )

    if check_type == "reference_exists":
        # source_column (str) 필수, target (str) 필수
        if not isinstance(rule_config.get("source_column"), str) or not rule_config.get("source_column"):
            raise HTTPException(
                status_code=422,
                detail="reference_exists 규칙은 'source_column' (문자열) 필드가 필수입니다",
            )
        if not isinstance(rule_config.get("target"), str) or not rule_config.get("target"):
            raise HTTPException(
                status_code=422,
                detail="reference_exists 규칙은 'target' (문자열) 필드가 필수입니다",
            )

    elif check_type == "compare_layers":
        # column (str), operator (str), reference_layer_column (str) 필수
        # threshold_ratio (number) 선택
        if not isinstance(rule_config.get("column"), str) or not rule_config.get("column"):
            raise HTTPException(
                status_code=422,
                detail="compare_layers 규칙은 'column' (문자열) 필드가 필수입니다",
            )
        operator = rule_config.get("operator")
        if operator not in ALLOWED_OPERATORS:
            raise HTTPException(
                status_code=422,
                detail=(
                    f"compare_layers 규칙의 operator는 다음 중 하나여야 합니다: "
                    f"{', '.join(ALLOWED_OPERATORS)}. 현재 값: '{operator}'"
                ),
            )
        if not isinstance(rule_config.get("reference_layer_column"), str) or not rule_config.get("reference_layer_column"):
            raise HTTPException(
                status_code=422,
                detail="compare_layers 규칙은 'reference_layer_column' (문자열) 필드가 필수입니다",
            )
        if "threshold_ratio" in rule_config:
            try:
                float(rule_config["threshold_ratio"])
            except (TypeError, ValueError):
                raise HTTPException(
                    status_code=422,
                    detail="compare_layers 규칙의 'threshold_ratio'는 숫자여야 합니다",
                )

    elif check_type == "equipment_compatibility":
        # column (str), compatibility (str) 필수
        # within_range 시 range_tolerance (number) 필수
        if not isinstance(rule_config.get("column"), str) or not rule_config.get("column"):
            raise HTTPException(
                status_code=422,
                detail="equipment_compatibility 규칙은 'column' (문자열) 필드가 필수입니다",
            )
        compatibility = rule_config.get("compatibility")
        if compatibility not in ALLOWED_COMPATIBILITY_TYPES:
            raise HTTPException(
                status_code=422,
                detail=(
                    f"equipment_compatibility 규칙의 compatibility는 다음 중 하나여야 합니다: "
                    f"{', '.join(ALLOWED_COMPATIBILITY_TYPES)}. 현재 값: '{compatibility}'"
                ),
            )
        if compatibility == "within_range":
            if "range_tolerance" not in rule_config:
                raise HTTPException(
                    status_code=422,
                    detail="within_range 호환성 타입은 'range_tolerance' (숫자) 필드가 필수입니다",
                )
            try:
                float(rule_config["range_tolerance"])
            except (TypeError, ValueError):
                raise HTTPException(
                    status_code=422,
                    detail="equipment_compatibility 규칙의 'range_tolerance'는 숫자여야 합니다",
                )


async def replace_validations(
    db: AsyncSession,
    column_id: int,
    rules: list[ValidationRuleCreate],
) -> dict:
    """Replace all validation rules for a column."""
    # Validate column exists
    col_def = await db.get(ColumnDefinition, column_id)
    if not col_def:
        raise HTTPException(status_code=404, detail="Column not found")

    # Validate all rules before making changes
    for rule in rules:
        # Validate rule_type
        if rule.rule_type not in ALLOWED_RULE_TYPES:
            raise HTTPException(status_code=422, detail=f"Invalid rule_type: {rule.rule_type}. Must be one of: {', '.join(ALLOWED_RULE_TYPES)}")

        # Validate rule_config for specific rule types
        if rule.rule_type == "range":
            if "min" not in rule.rule_config or "max" not in rule.rule_config:
                raise HTTPException(status_code=422, detail="Range rule requires 'min' and 'max' in rule_config")

        if rule.rule_type == "conditional_required":
            required_fields = ["condition_column", "condition_value", "operator"]
            missing_fields = [f for f in required_fields if f not in rule.rule_config]
            if missing_fields:
                raise HTTPException(
                    status_code=422,
                    detail=(
                        f"conditional_required rule requires: {', '.join(required_fields)}. "
                        f"Missing: {', '.join(missing_fields)}"
                    ),
                )

        if rule.rule_type == "cross_layer":
            _validate_cross_layer_rule_config(rule.rule_config)

    # Delete existing validations
    await db.execute(
        delete(ColumnValidation).where(ColumnValidation.column_id == column_id)
    )

    # Create new validations
    for rule in rules:
        validation = ColumnValidation(
            column_id=column_id,
            rule_type=rule.rule_type,
            rule_config=rule.rule_config,
            error_message=rule.error_message,
            is_active=rule.is_active,
        )
        db.add(validation)

    await db.commit()

    return {
        "column_id": column_id,
        "validation_count": len(rules),
        "message": f"Replaced {len(rules)} validation rule(s) for column {column_id}",
    }


async def bulk_upload_validations(
    db: AsyncSession,
    file: BinaryIO,
) -> BulkUploadResponse:
    """Bulk upload validation rules from Excel file."""
    try:
        wb = load_workbook(file)
        ws = wb.active
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to read Excel file: {str(e)}")

    # Read header row
    headers = [cell.value for cell in ws[1]]
    required_headers = ["column_name", "rule_type", "rule_config", "error_message", "is_active"]
    if not all(h in headers for h in required_headers):
        raise HTTPException(
            400,
            f"Excel file must have headers: {', '.join(required_headers)}"
        )

    # Map column names to column IDs
    result = await db.execute(select(ColumnDefinition))
    all_columns = result.scalars().all()
    column_map = {col.column_name: col for col in all_columns}

    # Process rows
    total_rows = 0
    columns_updated = set()
    rules_created = 0
    warnings = []

    # Group rules by column_name
    rules_by_column = {}

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not any(row):
            continue

        total_rows += 1

        # Parse row
        row_dict = dict(zip(headers, row))
        column_name = row_dict.get("column_name")
        rule_type = row_dict.get("rule_type")
        rule_config_str = row_dict.get("rule_config")
        error_message = row_dict.get("error_message")
        is_active_str = row_dict.get("is_active", "TRUE")

        # Validate column exists
        if column_name not in column_map:
            warnings.append(f"Row {row_idx}: Column '{column_name}' not found")
            continue

        # Validate rule_type
        if rule_type not in ALLOWED_RULE_TYPES:
            warnings.append(
                f"Row {row_idx}: Invalid rule_type '{rule_type}'. "
                f"Must be one of: {', '.join(ALLOWED_RULE_TYPES)}"
            )
            continue

        # Parse rule_config JSON
        try:
            rule_config = json.loads(rule_config_str) if rule_config_str else {}
        except json.JSONDecodeError:
            warnings.append(f"Row {row_idx}: Invalid JSON in rule_config: {rule_config_str}")
            continue

        # Parse is_active
        is_active = str(is_active_str).upper() in ("TRUE", "1", "YES", "Y")

        # Create ValidationRuleCreate object
        try:
            rule_create = ValidationRuleCreate(
                rule_type=rule_type,
                rule_config=rule_config,
                error_message=error_message,
                is_active=is_active,
            )
        except ValueError as e:
            warnings.append(f"Row {row_idx}: Validation error: {str(e)}")
            continue

        # Group by column
        if column_name not in rules_by_column:
            rules_by_column[column_name] = []
        rules_by_column[column_name].append(rule_create)

    # Replace validations for each column
    for column_name, rules in rules_by_column.items():
        col_def = column_map[column_name]
        try:
            await replace_validations(db, col_def.id, rules)
            columns_updated.add(column_name)
            rules_created += len(rules)
        except Exception as e:
            warnings.append(f"Column '{column_name}': Failed to replace validations: {str(e)}")

    return BulkUploadResponse(
        total_rows=total_rows,
        columns_updated=len(columns_updated),
        rules_created=rules_created,
        warnings=warnings,
    )


# ---------------------------------------------------------------------------
# Select Options Functions
# ---------------------------------------------------------------------------

async def list_select_columns(db: AsyncSession) -> list[ColumnSelectOptionsResponse]:
    """List columns with data_type='select'."""
    query = (
        select(ColumnDefinition, ColumnCategory)
        .join(ColumnCategory, ColumnDefinition.category_id == ColumnCategory.id)
        .where(ColumnDefinition.data_type == "select")
        .order_by(ColumnCategory.sort_order, ColumnDefinition.sort_order)
    )
    result = await db.execute(query)
    rows = result.all()

    return [
        ColumnSelectOptionsResponse(
            id=col.id,
            column_name=col.column_name,
            display_name=col.display_name,
            category_code=cat.category_code,
            data_type=col.data_type,
            select_options=col.select_options,
        )
        for col, cat in rows
    ]


async def update_select_options(
    db: AsyncSession, column_id: int, data: SelectOptionsUpdate
) -> ColumnSelectOptionsResponse:
    """Update select_options for a column."""
    col = await db.get(ColumnDefinition, column_id)
    if not col:
        raise HTTPException(status_code=404, detail="Column not found")
    if col.data_type != "select":
        raise HTTPException(status_code=400, detail="Column is not a select type")

    col.select_options = data.select_options
    await db.commit()
    await db.refresh(col)

    # Get category
    cat = await db.get(ColumnCategory, col.category_id)
    return ColumnSelectOptionsResponse(
        id=col.id,
        column_name=col.column_name,
        display_name=col.display_name,
        category_code=cat.category_code if cat else None,
        data_type=col.data_type,
        select_options=col.select_options,
    )


# ---------------------------------------------------------------------------
# Audit Log Functions
# ---------------------------------------------------------------------------

async def list_audit_logs(
    db: AsyncSession,
    project_id: int | None = None,
    line_id: int | None = None,
    changed_by: int | None = None,
    change_type: str | None = None,
    date_from: datetime | None = None,
    date_to: datetime | None = None,
    offset: int = 0,
    limit: int = 50,
) -> tuple[list[AuditLogEntry], int]:
    """List audit logs with filters. Returns (items, total_count)."""
    # Aliases for joins
    from sqlalchemy.orm import aliased
    from sqlalchemy import func

    UserAlias = aliased(User, name="changed_by_user")
    LayerAlias = aliased(Layer, name="layer_alias")

    base_query = (
        select(
            ChangeLog,
            Project.id.label("project_id"),
            Product.product_name.label("project_name"),
            LayerAlias.layer_name.label("layer_name"),
            UserAlias.display_name.label("changed_by_name"),
        )
        .outerjoin(ProjectLayer, ChangeLog.project_layer_id == ProjectLayer.id)
        .outerjoin(Project, ProjectLayer.project_id == Project.id)
        .outerjoin(Product, Project.product_id == Product.id)
        .outerjoin(LayerAlias, ProjectLayer.layer_id == LayerAlias.id)
        .outerjoin(UserAlias, ChangeLog.changed_by == UserAlias.id)
    )

    # Apply filters
    if project_id is not None:
        base_query = base_query.where(Project.id == project_id)
    if line_id is not None:
        base_query = base_query.where(Product.line_id == line_id)
    if changed_by is not None:
        base_query = base_query.where(ChangeLog.changed_by == changed_by)
    if change_type is not None:
        base_query = base_query.where(ChangeLog.change_type == change_type)
    if date_from is not None:
        base_query = base_query.where(ChangeLog.changed_at >= date_from)
    if date_to is not None:
        base_query = base_query.where(ChangeLog.changed_at <= date_to)

    # Count total
    count_query = select(func.count()).select_from(base_query.subquery())
    total_result = await db.execute(count_query)
    total = total_result.scalar_one()

    # Apply ordering and pagination
    paginated_query = (
        base_query
        .order_by(ChangeLog.changed_at.desc())
        .offset(offset)
        .limit(limit)
    )

    result = await db.execute(paginated_query)
    rows = result.all()

    items = [
        AuditLogEntry(
            id=row.ChangeLog.id,
            project_id=row.project_id,
            project_name=row.project_name,
            layer_name=row.layer_name,
            column_name=row.ChangeLog.column_name,
            old_value=row.ChangeLog.old_value,
            new_value=row.ChangeLog.new_value,
            change_type=row.ChangeLog.change_type,
            changed_by=row.ChangeLog.changed_by,
            changed_by_name=row.changed_by_name,
            changed_at=row.ChangeLog.changed_at,
        )
        for row in rows
    ]

    return items, total
