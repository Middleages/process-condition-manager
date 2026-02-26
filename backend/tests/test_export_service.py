"""
Tests for app.services.export_service.ExportService

This file covers ExportService orchestration (generate_type_a/b/c end-to-end).
Unit tests for pure builder functions (sanitize_filename, build_type_a_data,
build_type_b_data, build_type_c_data, write_excel) live in test_export_builders.py.
"""

import io
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.services.export_service import ExportService


# ---------------------------------------------------------------------------
# Helpers: lightweight mock objects that simulate ORM instances
# ---------------------------------------------------------------------------


def make_layer(layer_name: str) -> SimpleNamespace:
    """Create a mock Layer with a layer_name attribute."""
    return SimpleNamespace(layer_name=layer_name)


def make_col_def(column_name: str) -> SimpleNamespace:
    """Create a mock ColumnDefinition with a column_name attribute."""
    return SimpleNamespace(column_name=column_name)


def make_mapping(
    target_column_name: str,
    column_name: str,
    sort_order: int = 0,
    is_required: bool = False,
) -> SimpleNamespace:
    """Create a mock ExportColumnMapping."""
    return SimpleNamespace(
        target_column_name=target_column_name,
        column_definition=make_col_def(column_name),
        sort_order=sort_order,
        is_required=is_required,
        source_type="condition",
        data_source_id=None,
        source_column_name=None,
    )


def make_project_layer(
    layer_name: str,
    conditions: dict,
    pl_id: int = 1,
    sort_order: int = 0,
) -> SimpleNamespace:
    """Create a mock ProjectLayer with a nested layer."""
    return SimpleNamespace(
        id=pl_id,
        layer=make_layer(layer_name),
        conditions=conditions,
        sort_order=sort_order,
    )


# ---------------------------------------------------------------------------
# _generate_type_a, _generate_type_b, _generate_type_c (integration with _write_excel)
# ---------------------------------------------------------------------------


class TestGenerateTypes:
    """Verify that each _generate_type_X method produces valid xlsx bytes."""

    def setup_method(self):
        self.service = ExportService()

    def test_generate_type_a_produces_xlsx(self):
        layers = [make_project_layer("LAYER_A", {"speed": 2000}, pl_id=1)]
        mappings = [make_mapping("SPEED", "speed")]
        result = self.service._generate_type_a("PROD-A", layers, mappings)
        wb = load_workbook(io.BytesIO(result))
        ws = wb.active
        assert ws.cell(row=1, column=1).value == "LAYER_ID"
        assert ws.cell(row=2, column=1).value == "LAYER_A"

    def test_generate_type_b_produces_xlsx(self):
        """EQP column-based Type B Excel output produces valid xlsx."""
        conditions = {"speed": 2000, "EQP_01": "EQ-01"}
        layers = [make_project_layer("LAYER_A", conditions, pl_id=1)]
        mappings = [make_mapping("SPEED", "speed")]
        result = self.service._generate_type_b("PROD-A", layers, mappings)
        wb = load_workbook(io.BytesIO(result))
        ws = wb.active
        header_row = [ws.cell(row=1, column=i).value for i in range(1, 5)]
        assert "EQUIP_ID" in header_row

    def test_generate_type_c_produces_xlsx(self):
        layers = [make_project_layer("LAYER_A", {"speed": 2000}, pl_id=1)]
        mappings = [make_mapping("SPEED", "speed")]
        result = self.service._generate_type_c("PROD-A", layers, mappings, {"speed": "rpm"})
        wb = load_workbook(io.BytesIO(result))
        ws = wb.active
        header_row = [ws.cell(row=1, column=i).value for i in range(1, 6)]
        assert "PARAM_KEY" in header_row
        assert "PARAM_VALUE" in header_row
        assert "UNIT" in header_row
        # Verify unit appears in data
        data_row = [ws.cell(row=2, column=i).value for i in range(1, 6)]
        assert "rpm" in data_row
